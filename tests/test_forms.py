import json
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from backend.api.app import create_app
from backend.api.dependencies import get_rag_service
from backend.config import Settings, get_settings
from backend.forms import (
    FieldExtractor,
    FieldProposal,
    FormField,
    InvalidFormError,
    UnsupportedFormError,
    detect_fields,
    fill_form,
    filled_filename,
)
from backend.ingestion.scanner import FileMetadata
from backend.interfaces.chunker import DocumentChunk
from backend.interfaces.parser import ParsedDocument
from backend.interfaces.retriever import RetrievalResult

EXAMPLES = Path(__file__).resolve().parents[1] / "examples" / "forms"


def docx_bytes(build: Any) -> bytes:
    document = Document()
    build(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def docx_text(data: bytes) -> list[str]:
    return [paragraph.text for paragraph in Document(BytesIO(data)).paragraphs if paragraph.text]


def pdf_form_bytes() -> bytes:
    output = BytesIO()
    page = canvas.Canvas(output, pagesize=letter)
    page.drawString(72, 720, "Form")
    page.acroForm.textfield(
        name="full_name", tooltip="Full name", x=72, y=680, width=200, height=20
    )
    page.acroForm.textfield(name="homePhone", x=72, y=640, width=200, height=20)
    page.acroForm.checkbox(name="agree", tooltip="I agree", x=72, y=600, size=14)
    page.save()
    return output.getvalue()


# --- Word forms ------------------------------------------------------------


def test_word_fields_are_found_in_lines_tables_and_placeholders() -> None:
    def build(document: Any) -> None:
        document.add_paragraph("Please fill this in. Full name: ____________")
        document.add_paragraph("Email: ________  Phone: ........")
        document.add_paragraph("Emergency contact:")
        document.add_paragraph("Hello {{Company}}, welcome {{Company}}")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Employer:"
        table.rows[1].cells[0].text = "Title"
        table.rows[1].cells[1].text = "already filled"

    fields = detect_fields("form.docx", docx_bytes(build))

    assert [(field.id, field.label) for field in fields] == [
        ("p0.0", "Full name"),
        ("p1.0", "Email"),
        ("p1.1", "Phone"),
        ("p2", "Emergency contact"),
        ("ph:Company", "Company"),
        ("t0r0c1", "Employer"),
    ]


def test_ordinary_sentences_are_not_mistaken_for_fields() -> None:
    def build(document: Any) -> None:
        document.add_paragraph("Note: bring a photo ID.")
        document.add_paragraph("Total: 42")

    assert detect_fields("form.docx", docx_bytes(build)) == []


def test_word_fill_replaces_blanks_and_keeps_unfilled_lines() -> None:
    def build(document: Any) -> None:
        document.add_paragraph("Name: ______________")
        document.add_paragraph("Email: ______  Phone: ______")
        document.add_paragraph("Contact:")
        document.add_paragraph("Hi {{Company}}!")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Employer"

    filled = fill_form(
        "form.docx",
        docx_bytes(build),
        {
            "p0.0": "Alex Rivera",
            "p1.0": "alex@example.com",
            "p2": "Jordan",
            "ph:Company": "Northwind",
            "t0r0c1": "Northwind Analytics",
        },
    )

    assert docx_text(filled)[:4] == [
        "Name: Alex Rivera",
        "Email: alex@example.com  Phone: ______",
        "Contact: Jordan",
        "Hi Northwind!",
    ]
    assert Document(BytesIO(filled)).tables[0].rows[0].cells[1].text == "Northwind Analytics"


def test_word_fill_handles_blanks_split_across_runs() -> None:
    def build(document: Any) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run("Name: __")
        paragraph.add_run("____").bold = True

    filled = fill_form("form.docx", docx_bytes(build), {"p0.0": "Alex"})

    assert docx_text(filled) == ["Name: Alex"]


def test_blank_values_leave_the_form_unchanged() -> None:
    def build(document: Any) -> None:
        document.add_paragraph("Name: ________")

    filled = fill_form("form.docx", docx_bytes(build), {"p0.0": "   "})

    assert docx_text(filled) == ["Name: ________"]


def test_unreadable_word_files_are_rejected() -> None:
    with pytest.raises(InvalidFormError):
        detect_fields("form.docx", b"not a zip")


# --- PDF forms -------------------------------------------------------------


def test_pdf_fields_use_tooltips_and_readable_names() -> None:
    fields = detect_fields("form.pdf", pdf_form_bytes())

    assert [(field.id, field.label, field.kind) for field in fields] == [
        ("full_name", "Full name", "text"),
        ("homePhone", "Home phone", "text"),
        ("agree", "I agree", "checkbox"),
    ]


def test_pdf_fill_sets_text_and_ticks_checkboxes() -> None:
    filled = fill_form(
        "form.pdf",
        pdf_form_bytes(),
        {"full_name": "Alex Rivera", "agree": "yes", "homePhone": "", "unknown": "x"},
    )

    values = {
        name: field.get("/V") for name, field in PdfReader(BytesIO(filled)).get_fields().items()
    }
    assert values["full_name"] == "Alex Rivera"
    assert values["agree"] == "/Yes"
    assert values["homePhone"] in (None, "")


def test_pdf_checkbox_stays_off_unless_asked() -> None:
    filled = fill_form("form.pdf", pdf_form_bytes(), {"agree": "no"})

    value = PdfReader(BytesIO(filled)).get_fields()["agree"].get("/V")
    assert value in (None, "", "/Off")


def test_unreadable_pdfs_are_rejected() -> None:
    with pytest.raises(InvalidFormError):
        detect_fields("form.pdf", b"%PDF-broken")


def test_other_file_types_are_unsupported() -> None:
    with pytest.raises(UnsupportedFormError):
        detect_fields("form.txt", b"hello")


def test_filled_files_get_a_clear_name() -> None:
    assert filled_filename("My Form.DOCX") == "My Form-filled.docx"


@pytest.mark.parametrize(
    ("name", "labels"),
    [
        (
            "job_application.docx",
            {"Full name", "Email", "Phone", "Date of birth", "Current employer", "Job title"},
        ),
        ("event_registration.docx", {"Full name", "Email", "Phone", "Dietary restrictions"}),
        ("membership_form.pdf", {"Full name", "Email address", "Home address", "Date of birth"}),
    ],
)
def test_bundled_example_forms_are_detected(name: str, labels: set[str]) -> None:
    fields = detect_fields(name, (EXAMPLES / name).read_bytes())

    assert labels <= {field.label for field in fields}


# --- proposing values ------------------------------------------------------


def passage(content: str, name: str = "profile.txt") -> RetrievalResult:
    path = Path(name)
    document = ParsedDocument(
        source_path=path,
        file_type=".txt",
        content=content,
        metadata=FileMetadata(name=name, path=path, extension=".txt", size_bytes=len(content)),
    )
    chunk = DocumentChunk(
        source_document=document,
        content=content,
        chunk_index=0,
        start_char=0,
        end_char=len(content),
    )
    return RetrievalResult(chunk=chunk, score=1.0, retriever="test")


class ScriptedGenerator:
    def __init__(self, replies: dict[str, str | Exception]) -> None:
        self.replies = replies
        self.prompts: list[str] = []

    def generate(self, prompt: str) -> str:
        self.prompts.append(prompt)
        for label, reply in self.replies.items():
            if f"FIELD\n{label}\n" in prompt:
                if isinstance(reply, Exception):
                    raise reply
                return reply
        return "NOT FOUND"


PROFILE = passage(
    "Email: alex.rivera@example.com\nPhone: (555) 014-2290\nDate of birth: March 14, 1991"
)


def extractor(
    replies: dict[str, str | Exception], found: Sequence[RetrievalResult] = (PROFILE,)
) -> FieldExtractor:
    return FieldExtractor(retrieve=lambda query: list(found), generator=ScriptedGenerator(replies))


def test_a_value_found_word_for_word_is_verified_with_its_source() -> None:
    proposal = extractor({"Email": "alex.rivera@example.com"}).propose(FormField("e", "Email"))

    assert proposal.value == "alex.rivera@example.com"
    assert proposal.verified
    assert proposal.source == "profile.txt"
    assert proposal.snippet is not None and "alex.rivera@example.com" in proposal.snippet


def test_formatting_differences_do_not_hide_a_match() -> None:
    proposal = extractor({"Phone": '"555 014 2290"'}).propose(FormField("p", "Phone"))

    assert proposal.value == "555 014 2290"
    assert proposal.verified


def test_a_value_that_is_not_in_the_documents_is_flagged_unverified() -> None:
    proposal = extractor({"Email": "someone@else.com"}).propose(FormField("e", "Email"))

    assert proposal.value == "someone@else.com"
    assert not proposal.verified
    assert proposal.source is None


@pytest.mark.parametrize(
    "reply",
    [
        "NOT FOUND",
        "not found.",
        "I don't know",
        "N/A",
        "",
        "   \n",
        "x" * 400,
        RuntimeError("model down"),
    ],
)
def test_missing_or_unusable_answers_leave_the_field_blank(reply: str | Exception) -> None:
    proposal = extractor({"Salary": reply}).propose(FormField("s", "Salary"))

    assert proposal.value is None
    assert not proposal.verified


def test_no_retrieved_passages_means_no_proposal_and_no_model_call() -> None:
    generator = ScriptedGenerator({"Email": "x"})
    result = FieldExtractor(retrieve=lambda query: [], generator=generator).propose(
        FormField("e", "Email")
    )

    assert result == FieldProposal(field=FormField("e", "Email"))
    assert generator.prompts == []


def test_checkboxes_are_left_for_the_user() -> None:
    generator = ScriptedGenerator({})
    result = FieldExtractor(retrieve=lambda query: [PROFILE], generator=generator).propose(
        FormField("a", "I agree", kind="checkbox")
    )

    assert result.value is None
    assert generator.prompts == []


def test_prompt_forbids_guessing_and_quotes_only_retrieved_text() -> None:
    generator = ScriptedGenerator({"Email": "alex.rivera@example.com"})
    FieldExtractor(retrieve=lambda query: [PROFILE], generator=generator).propose(
        FormField("e", "Email")
    )

    prompt = generator.prompts[0]
    assert "Do not guess" in prompt
    assert "NOT FOUND" in prompt
    assert "alex.rivera@example.com" in prompt


# --- API -------------------------------------------------------------------


class FakeFormService:
    def propose_form_values(self, fields: Sequence[FormField]) -> list[FieldProposal]:
        return [
            FieldProposal(
                field=field,
                value="Alex Rivera" if field.label == "Name" else None,
                verified=field.label == "Name",
                source="profile.txt" if field.label == "Name" else None,
            )
            for field in fields
        ]


@pytest.fixture
def client() -> TestClient:
    app = create_app()
    app.dependency_overrides[get_rag_service] = lambda: FakeFormService()
    return TestClient(app)


def name_form() -> bytes:
    return docx_bytes(lambda document: document.add_paragraph("Name: ________"))


def test_propose_returns_fields_with_suggestions(client: TestClient) -> None:
    response = client.post(
        "/forms/propose",
        files={"file": ("form.docx", name_form(), "application/octet-stream")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "form.docx"
    assert body["fields"] == [
        {
            "id": "p0.0",
            "label": "Name",
            "kind": "text",
            "value": "Alex Rivera",
            "verified": True,
            "source": "profile.txt",
            "snippet": None,
        }
    ]


@pytest.mark.parametrize(
    ("filename", "content", "status"),
    [
        ("notes.txt", b"hello", 415),
        ("empty.docx", b"", 422),
        ("broken.docx", b"not a zip", 422),
        ("plain.docx", docx_bytes(lambda document: document.add_paragraph("Just text.")), 422),
    ],
    ids=["wrong-type", "empty", "corrupt", "no-fields"],
)
def test_propose_rejects_files_that_are_not_usable_forms(
    client: TestClient, filename: str, content: bytes, status: int
) -> None:
    response = client.post(
        "/forms/propose", files={"file": (filename, content, "application/octet-stream")}
    )

    assert response.status_code == status


def test_propose_rejects_oversized_forms() -> None:
    app = create_app()
    app.dependency_overrides[get_rag_service] = lambda: FakeFormService()
    app.dependency_overrides[get_settings] = lambda: Settings(max_upload_bytes=100)

    response = TestClient(app).post(
        "/forms/propose",
        files={"file": ("form.docx", name_form(), "application/octet-stream")},
    )

    assert response.status_code == 413


def test_fill_returns_the_completed_document(client: TestClient) -> None:
    response = client.post(
        "/forms/fill",
        files={"file": ("form.docx", name_form(), "application/octet-stream")},
        data={"values": json.dumps({"p0.0": "Alex Rivera"})},
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="form-filled.docx"'
    assert docx_text(response.content) == ["Name: Alex Rivera"]


@pytest.mark.parametrize("values", ["not json", "[1, 2]", '"text"'])
def test_fill_requires_a_json_object(client: TestClient, values: str) -> None:
    response = client.post(
        "/forms/fill",
        files={"file": ("form.docx", name_form(), "application/octet-stream")},
        data={"values": values},
    )

    assert response.status_code == 422
