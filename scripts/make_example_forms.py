"""
Create the example forms and the matching profile document in examples/forms/.

    uv run python scripts/make_example_forms.py

Everything here is fictional. To try the form filler, ingest examples/forms/sample_profile.txt
into the app, then upload one of the forms on the Forms tab.
"""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

OUTPUT = Path(__file__).resolve().parent.parent / "examples" / "forms"

PROFILE = """\
Personal profile: Alex Rivera

Full name: Alex Rivera
Date of birth: March 14, 1991
Email: alex.rivera@example.com
Phone: (555) 014-2290
Home address: 482 Maple Street, Springfield, IL 62704

Current employer: Northwind Analytics
Job title: Data Engineer
Years of experience: 7

Emergency contact: Jordan Rivera, (555) 014-7731
Preferred contact method: email
"""


def job_application() -> None:
    document = Document()
    document.add_heading("Job application", level=1)
    document.add_paragraph("Please complete every line you can.")
    document.add_paragraph("Full name: ______________________________")
    document.add_paragraph("Email: ____________________   Phone: ____________________")
    document.add_paragraph("Date of birth: ____________")

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, label in zip(
        table.rows, ("Current employer", "Job title", "Years of experience"), strict=True
    ):
        row.cells[0].text = label

    document.add_paragraph("Salary expectations: ____________")
    document.add_paragraph("Signature: ____________________")
    document.save(OUTPUT / "job_application.docx")


def event_registration() -> None:
    document = Document()
    document.add_heading("Event registration", level=1)
    document.add_paragraph("Name: {{Full name}}")
    document.add_paragraph("Contact email: {{Email}}")
    document.add_paragraph("Phone: {{Phone}}")
    document.add_paragraph("Emergency contact:")
    document.add_paragraph("Dietary restrictions: {{Dietary restrictions}}")
    document.save(OUTPUT / "event_registration.docx")


def membership_form() -> None:
    page = canvas.Canvas(str(OUTPUT / "membership_form.pdf"), pagesize=letter)
    page.setFont("Helvetica-Bold", 16)
    page.drawString(72, 720, "Community Gym membership")
    page.setFont("Helvetica", 11)

    rows = (
        ("full_name", "Full name"),
        ("email", "Email address"),
        ("phone", "Phone number"),
        ("address", "Home address"),
        ("date_of_birth", "Date of birth"),
    )
    y = 670
    for name, label in rows:
        page.drawString(72, y + 4, f"{label}:")
        page.acroForm.textfield(
            name=name,
            tooltip=label,
            x=210,
            y=y - 4,
            width=300,
            height=22,
            borderStyle="solid",
            forceBorder=True,
        )
        y -= 40

    page.drawString(72, y + 4, "Send me the newsletter")
    page.acroForm.checkbox(name="newsletter", tooltip="Newsletter", x=210, y=y - 4, size=16)
    page.save()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    (OUTPUT / "sample_profile.txt").write_text(PROFILE, encoding="utf-8")
    job_application()
    event_registration()
    membership_form()
    print(f"Wrote examples to {OUTPUT}")


if __name__ == "__main__":
    main()
