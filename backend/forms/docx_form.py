import re
from collections.abc import Iterator, Mapping
from io import BytesIO
from typing import Any

from docx import Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from backend.forms.models import FormField, InvalidFormError

_BLANK = r"(?:_{3,}|\.{4,}|…{2,})"
_LABEL_CHARS = r"[A-Za-z][A-Za-z0-9 /'’()&-]{0,50}?"
_INLINE = re.compile(rf"({_LABEL_CHARS})\s*:\s*({_BLANK})")
_TRAILING_LABEL = re.compile(rf"^\s*({_LABEL_CHARS})\s*:\s*$")
_PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_BLANK_ONLY = re.compile(rf"^\s*{_BLANK}\s*$")


def _open(data: bytes) -> Any:
    try:
        return Document(BytesIO(data))
    except Exception as exc:
        raise InvalidFormError("The Word document could not be read.") from exc


def _unique_cells(row: Any) -> list[tuple[int, _Cell]]:
    """Cells of a row with merged cells counted once, keeping their column index."""

    unique: list[tuple[int, _Cell]] = []
    for column, cell in enumerate(row.cells):
        if not any(cell._tc is seen._tc for _, seen in unique):
            unique.append((column, cell))
    return unique


# Ids starting with "p" are body paragraphs; table cell paragraphs start with "t".
def _paragraph_sites(document: Any) -> Iterator[tuple[str, Paragraph]]:
    for index, paragraph in enumerate(document.paragraphs):
        yield f"p{index}", paragraph

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column, cell in _unique_cells(row):
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield f"t{table_index}r{row_index}c{column}p{paragraph_index}", paragraph


def _cell_sites(document: Any) -> Iterator[tuple[str, str, _Cell]]:
    """Empty cells that sit to the right of a cell holding a label."""

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            cells = _unique_cells(row)
            for (_, label_cell), (column, value_cell) in zip(cells, cells[1:], strict=False):
                label = label_cell.text.strip().rstrip(":").strip()
                value = value_cell.text
                if (
                    label
                    and not _BLANK_ONLY.match(label)
                    and (not value.strip() or _BLANK_ONLY.match(value))
                ):
                    yield f"t{table_index}r{row_index}c{column}", label, value_cell


def detect_docx_fields(data: bytes) -> list[FormField]:
    document = _open(data)
    fields: list[FormField] = []
    placeholders: set[str] = set()

    for prefix, paragraph in _paragraph_sites(document):
        text = paragraph.text
        inline = list(_INLINE.finditer(text))

        for position, match in enumerate(inline):
            fields.append(FormField(id=f"{prefix}.{position}", label=match.group(1).strip()))

        if not inline and prefix.startswith("p") and (trailing := _TRAILING_LABEL.match(text)):
            fields.append(FormField(id=prefix, label=trailing.group(1).strip()))

        for placeholder in _PLACEHOLDER.finditer(text):
            label = placeholder.group(1).strip()
            if label not in placeholders:
                placeholders.add(label)
                fields.append(FormField(id=f"ph:{label}", label=label))

    for field_id, label, _ in _cell_sites(document):
        fields.append(FormField(id=field_id, label=label))

    return fields


def _replace_span(paragraph: Paragraph, start: int, end: int, new_text: str) -> None:
    """Replace characters [start, end) of the paragraph, keeping the surrounding formatting."""

    position = 0
    replaced = False
    for run in paragraph.runs:
        text = run.text
        run_start, run_end = position, position + len(text)
        position = run_end
        if run_end <= start or run_start >= end:
            continue

        local_start = max(start, run_start) - run_start
        local_end = min(end, run_end) - run_start
        run.text = text[:local_start] + ("" if replaced else new_text) + text[local_end:]
        replaced = True


def _set_cell(cell: _Cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _fill_paragraph(prefix: str, paragraph: Paragraph, values: Mapping[str, str]) -> None:
    inline = list(_INLINE.finditer(paragraph.text))
    for position in reversed(range(len(inline))):
        value = values.get(f"{prefix}.{position}", "").strip()
        if value:
            _replace_span(paragraph, inline[position].start(2), inline[position].end(2), value)

    if not inline and prefix.startswith("p") and _TRAILING_LABEL.match(paragraph.text):
        value = values.get(prefix, "").strip()
        if value:
            paragraph.add_run(f" {value}")

    for placeholder in reversed(list(_PLACEHOLDER.finditer(paragraph.text))):
        value = values.get(f"ph:{placeholder.group(1).strip()}", "").strip()
        if value:
            _replace_span(paragraph, placeholder.start(), placeholder.end(), value)


def fill_docx(data: bytes, values: Mapping[str, str]) -> bytes:
    document = _open(data)

    for prefix, paragraph in _paragraph_sites(document):
        _fill_paragraph(prefix, paragraph, values)

    for field_id, _, cell in _cell_sites(document):
        value = values.get(field_id, "").strip()
        if value:
            _set_cell(cell, value)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
